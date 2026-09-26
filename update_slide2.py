import sys
from pptx import Presentation

pptx_path = r'c:\Users\sricharan\Desktop\Project Samanvaya\Team_Panchajanyam_SIH26047_Samanvaya.pptx'
prs = Presentation(pptx_path)
slide2 = prs.slides[1]

replacements = {
    7: "A zero-hardware-footprint, dual-branch AI platform that operates in waiting rooms to digitize vernacular history taking, decipher handwritten prescriptions, and generate ABDM-compliant EHR case sheets in under 20 seconds.",
    8: "CORE SYSTEM CAPABILITIES",
    11: "Dual-System History Intake",
    12: "Adaptive voice/touch intake in 22 Indic languages via Bhashini AI; includes AIIA-aligned Dashavidha Pariksha for Ayush OPDs.",
    15: "Multimodal Prescription AI",
    16: "Deciphers handwritten doctor slips, lab reports and OPD cards via a 4-tier OCR cascade (Nemotron + Llama 3.2 Vision + Groq 120B).",
    19: "Instant Physician Case Sheet",
    20: "Synthesizes patient inputs into a structured draft EHR summary (Chief Complaint → HPI → ROS) for 1-click doctor verification.",
    23: "DPDP Act Vault & ABDM FHIR",
    24: "Statutory DPDP Act 2023 digital consent modal; session-purged storage; direct ABHA FHIR R4 JSON push to hospital HIS.",
    26: "GAME-CHANGING INNOVATIONS",
    29: "Govt Scheme Evaluation Engine",
    30: "Auto-evaluates PM-JAY, AB-PMJAY Senior Citizen (70+), Tele-MANAS & State Scheme eligibility directly from intake diagnosis.",
    33: "Floating AI Health Coach",
    34: "Interactive vernacular voice & chat assistant providing post-consultation prescription explanation and lifestyle guidance.",
    37: "PM Jan Aushadhi Savings Engine",
    38: "Matches brand drugs to generic PMBJP equivalents with 80%+ cost savings & pinpoints nearest Kendras via live GPS.",
    41: "WHO AWARE Stewardship Audit",
    42: "Real-time clinical safety net auditing prescribed antibiotics against WHO Access, Watch, Reserve tiers to combat AMR."
}

for shape_idx, text in replacements.items():
    if shape_idx < len(slide2.shapes):
        shape = slide2.shapes[shape_idx]
        if shape.has_text_frame:
            shape.text_frame.text = text

prs.save(pptx_path)
print("Successfully updated Slide 2 of Team_Panchajanyam_SIH26047_Samanvaya.pptx!")
